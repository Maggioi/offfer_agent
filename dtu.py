import docx
import copy

def duplicate_table_underneath(template_table, target_position_table, doc):
    """
    Creates a copy of the 'template_table' (which holds the clean placeholders)
    and inserts it directly below the 'target_position_table' (the current last table).
    """
    # 1. Create a native paragraph to guarantee separation
    blank_p = doc.add_paragraph()
    blank_p.add_run("\u00A0") 
    
    p_element = blank_p._p
    
    # WE DEEPCOPY THE CLEAN TEMPLATE, NOT THE MODIFIED ONE
    tbl_element = template_table._tbl
    new_tbl_element = copy.deepcopy(tbl_element)
    
    # We find the index of the CURRENT LAST TABLE to insert the new items below it
    target_tbl_element = target_position_table._tbl
    parent_element = target_tbl_element.getparent()
    current_index = parent_element.index(target_tbl_element)
    
    # 2. Move the paragraph separator below the current last table
    parent_element.remove(p_element)
    parent_element.insert(current_index + 1, p_element)
    
    # 3. Insert the clean duplicated table below the separator
    parent_element.insert(current_index + 2, new_tbl_element)
    
    # Wrap the new XML element back into a python-docx Table object
    return docx.table.Table(new_tbl_element, template_table._parent)